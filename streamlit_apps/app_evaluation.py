"""
5-Day Live Evaluation  —  Plan → Execute → Track

Create a session (tickers + portfolio + start date), then fill in 5 day-slots.
Each slot: run the pipeline or log manually, then check the T+1 outcome.
Everything is deletable. KB learns from every checked outcome.

  streamlit run streamlit_apps/app_evaluation.py
"""

import streamlit as st
import json, os, time
import pandas as pd
from datetime import datetime, date, timedelta
from typing import Dict, List, Optional
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from preferences import (
    SUPPORTED_TICKERS,
    PRESET_PORTFOLIOS,
    DEFAULT_PORTFOLIO_KEY,
)
from utils.portfolio_utils import _fetch_current_price

# ═══════════════════════════════════════════════════════════
# Paths
# ═══════════════════════════════════════════════════════════
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_PATH = os.path.join(ROOT, "results", "_eval_sessions.json")
EXCEL_PATH = os.path.join(ROOT, "results", "_eval_results.xlsx")
os.makedirs(os.path.join(ROOT, "results"), exist_ok=True)


# ═══════════════════════════════════════════════════════════
# Cached heavy object
# ═══════════════════════════════════════════════════════════
@st.cache_resource
def _graph():
    from graphs.trading_graph import TradingGraph
    return TradingGraph()


# ═══════════════════════════════════════════════════════════
# Helpers
# ═══════════════════════════════════════════════════════════
def _bdays(start: date, n: int) -> List[date]:
    """Next *n* business days starting from *start* (inclusive if weekday)."""
    out, cur = [], start
    while len(out) < n:
        if cur.weekday() < 5:
            out.append(cur)
        cur += timedelta(days=1)
    return out


def _next_bday(d: date) -> date:
    n = d + timedelta(days=1)
    while n.weekday() >= 5:
        n += timedelta(days=1)
    return n


def _price(ticker: str) -> Optional[float]:
    """Live (or last-close) price via yfinance."""
    try:
        p, _ = _fetch_current_price(ticker, 0.0)
        return round(p, 2) if p > 0 else None
    except Exception:
        return None


def _eval(decision: str, ret: float) -> bool:
    """Same rule as Explainability._evaluate."""
    d = decision.upper()
    if d == "BUY":
        return ret > 0
    if d == "SELL":
        return ret < 0
    return abs(ret) < 0.05


def _make_entry(price: float, result: dict, source: str = "pipeline") -> dict:
    """Build entry dict from pipeline result + price. Reused by run & re-run."""
    return {
        "entry_price": price,
        "decision": result.get("decision", "HOLD"),
        "confidence": result.get("confidence", 0.5),
        "reasoning_preview": result.get("reasoning", "")[:300],
        "logged_at": datetime.now().isoformat(),
        "source": source,
        "t1_price": None,
        "t1_return_pct": None,
        "was_correct": None,
        "outcome_checked_at": None,
        "notes": "",
    }


def _ts_label(iso: str) -> str:
    """Format ISO timestamp to short time label, e.g. '09:32 AM'."""
    try:
        return datetime.fromisoformat(iso).strftime("%I:%M %p")
    except Exception:
        return ""


# ═══════════════════════════════════════════════════════════
# Data I/O
# ═══════════════════════════════════════════════════════════
def _load() -> dict:
    if os.path.exists(DATA_PATH):
        try:
            with open(DATA_PATH) as f:
                return json.load(f)
        except (json.JSONDecodeError, IOError):
            pass
    return {"active": None, "sessions": {}}


def _save(data: dict):
    with open(DATA_PATH, "w") as f:
        json.dump(data, f, indent=2, default=str)
    _xlsx(data)


def _flat(data: dict, session_name: str = None) -> list:
    """Flatten entries → [{session, day, ticker, …}, …]."""
    out = []
    for sn, sess in data.get("sessions", {}).items():
        if session_name and sn != session_name:
            continue
        for dk, day in sess.get("days", {}).items():
            for t, e in day.get("entries", {}).items():
                if e and e.get("decision"):
                    out.append({
                        **e, "session": sn, "day": int(dk),
                        "ticker": t,
                        "log_date": day.get("log_date", ""),
                        "check_date": day.get("check_date", ""),
                    })
    return out


def _xlsx(data: dict):
    """Auto-sync Excel whenever JSON is saved."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Log"
    hdr = [
        "Session", "Day", "Log Date", "T+1 Date", "Ticker", "Decision",
        "Conf", "Entry $", "Logged At", "T+1 Price $", "Return", "Correct",
    ]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
    gf = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    rf = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for e in _flat(data):
        ret = e.get("t1_return_pct")
        ws.append([
            e["session"], e["day"], e["log_date"], e.get("check_date", ""),
            e["ticker"],
            e.get("decision", ""), f"{e.get('confidence', 0):.0%}",
            e.get("entry_price", ""),
            _ts_label(e.get("logged_at", "")),
            e.get("t1_price", ""),
            f"{ret * 100:+.2f}%" if ret is not None else "",
            {True: "YES", False: "NO"}.get(e.get("was_correct"), ""),
        ])
        if e.get("was_correct") is not None:
            fill = gf if e["was_correct"] else rf
            for col in range(1, len(hdr) + 1):
                ws.cell(row=ws.max_row, column=col).fill = fill

    for i in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    # Summary sheet
    all_entries = _flat(data)
    checked = [e for e in all_entries if e.get("was_correct") is not None]
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)

    n_total = len(all_entries)
    n_chk = len(checked)
    n_ok = sum(1 for e in checked if e["was_correct"])
    win_rate = n_ok / n_chk if n_chk else None
    avg_ret = sum(e.get("t1_return_pct", 0) for e in checked) / n_chk if n_chk else None

    early_s = [e for e in checked if e["day"] <= 2]
    late_s = [e for e in checked if e["day"] >= 4]
    ew_s = sum(1 for e in early_s if e["was_correct"]) / len(early_s) if early_s else None
    lw_s = sum(1 for e in late_s if e["was_correct"]) / len(late_s) if late_s else None

    always_buy = sum(1 for e in checked if e.get("t1_return_pct", 0) > 0)
    ab_rate = always_buy / n_chk if n_chk else None

    buys = sum(1 for e in all_entries if e.get("decision") == "BUY")
    sells = sum(1 for e in all_entries if e.get("decision") == "SELL")
    holds = sum(1 for e in all_entries if e.get("decision") == "HOLD")

    summary_rows = [
        ("Total logged", n_total),
        ("Total checked", n_chk),
        ("Win rate", f"{win_rate:.0%}" if win_rate is not None else "—"),
        ("Avg return", f"{avg_ret * 100:+.2f}%" if avg_ret is not None else "—"),
        ("Day 1-2 win rate", f"{ew_s:.0%}" if ew_s is not None else "—"),
        ("Day 4-5 win rate", f"{lw_s:.0%}" if lw_s is not None else "—"),
        ("KB delta", f"{lw_s - ew_s:+.0%}" if (ew_s is not None and lw_s is not None) else "—"),
        ("Always-BUY baseline", f"{ab_rate:.0%}" if ab_rate is not None else "—"),
        ("BUY decisions", buys),
        ("SELL decisions", sells),
        ("HOLD decisions", holds),
    ]
    tickers_seen = sorted({e["ticker"] for e in all_entries})
    for tk in tickers_seen:
        te = [e for e in checked if e["ticker"] == tk]
        if te:
            tw = sum(1 for e in te if e["was_correct"]) / len(te)
            tr = sum(e.get("t1_return_pct", 0) for e in te) / len(te)
            summary_rows.append((f"{tk} win rate", f"{tw:.0%}"))
            summary_rows.append((f"{tk} avg return", f"{tr * 100:+.2f}%"))

    for r in summary_rows:
        ws2.append(r)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 14

    wb.save(EXCEL_PATH)


# ═══════════════════════════════════════════════════════════
# Page config
# ═══════════════════════════════════════════════════════════
st.set_page_config(page_title="5-Day Eval", page_icon="📊", layout="wide")
st.markdown(
    "<style>.block-container{padding-top:.6rem}"
    " div[data-testid='stMetric']{text-align:center}</style>",
    unsafe_allow_html=True,
)

data = _load()
TODAY = date.today()


# ═══════════════════════════════════════════════════════════
# SIDEBAR — session management
# ═══════════════════════════════════════════════════════════
with st.sidebar:
    st.header("📋 Sessions")

    names = list(data.get("sessions", {}).keys())

    # ── select active ──
    if names:
        act = data.get("active") or names[0]
        if act not in names:
            act = names[0]
        data["active"] = st.selectbox("Active session", names, index=names.index(act))

    # ── new session ──
    st.divider()
    with st.expander("➕ New session", expanded=not names):
        nn = st.text_input("Name", value=f"Round {len(names) + 1}", key="ns_n")
        nt = st.multiselect(
            "Tickers (fixed for all 5 days)",
            SUPPORTED_TICKERS,
            default=["AAPL", "NVDA"],
            key="ns_t",
        )
        pk = [k for k in PRESET_PORTFOLIOS if k != "custom"]
        pl = [PRESET_PORTFOLIOS[k]["label"] for k in pk]
        di = pk.index(DEFAULT_PORTFOLIO_KEY) if DEFAULT_PORTFOLIO_KEY in pk else 0
        sl = st.selectbox("Portfolio", pl, index=di, key="ns_p")
        sk = pk[pl.index(sl)]
        sd = st.date_input("Start date", value=TODAY, key="ns_d")
        sn_notes = st.text_area(
            "Session notes (optional)",
            placeholder="e.g. KB empty, testing baseline…",
            height=68, key="ns_notes",
        )

        ok = bool(nn and nt and nn not in names)
        if st.button("✅ Create plan", type="primary", use_container_width=True, disabled=not ok):
            ld = _bdays(sd, 5)
            cd = [_next_bday(d) for d in ld]
            sess = {
                "tickers": nt,
                "portfolio_key": sk,
                "portfolio_label": sl,
                "portfolio": PRESET_PORTFOLIOS[sk]["holdings"],
                "created_at": datetime.now().isoformat(),
                "notes": sn_notes.strip(),
                "days": {},
            }
            for i, (l, c) in enumerate(zip(ld, cd), 1):
                sess["days"][str(i)] = {
                    "log_date": l.isoformat(),
                    "check_date": c.isoformat(),
                    "entries": {t: {} for t in nt},
                }
            data["sessions"][nn] = sess
            data["active"] = nn
            _save(data)
            st.rerun()

    # ── danger zone ──
    if names:
        st.divider()
        with st.expander("🗑️ Danger zone"):
            td = st.selectbox("Delete session", names, key="ds")
            if st.button("Delete forever", type="secondary", use_container_width=True):
                del data["sessions"][td]
                rem = list(data["sessions"].keys())
                data["active"] = rem[0] if rem else None
                _save(data)
                st.rerun()


# ═══════════════════════════════════════════════════════════
# Gate: need an active session
# ═══════════════════════════════════════════════════════════
act_name = data.get("active")
if not act_name or act_name not in data.get("sessions", {}):
    st.title("📊 5-Day Evaluation")
    st.info("👈 Create a session in the sidebar to start.")
    st.stop()

sess = data["sessions"][act_name]
tickers = sess["tickers"]
portfolio = sess.get("portfolio", [])


# ═══════════════════════════════════════════════════════════
# Title row + metrics
# ═══════════════════════════════════════════════════════════
st.markdown(f"## 📊 {act_name}")
_cap = [
    f"**{sess.get('portfolio_label', '')}**",
    ", ".join(tickers),
    f"started {sess.get('created_at', '')[:10]}",
]
if sess.get("notes"):
    _cap.append(f"*{sess['notes']}*")
st.caption("  ·  ".join(_cap))

se = _flat(data, act_name)
se_chk = [e for e in se if e.get("was_correct") is not None]
se_ok = [e for e in se_chk if e["was_correct"]]
se_pend = sum(1 for e in se if e.get("decision") and e.get("t1_price") is None)

m1, m2, m3, m4 = st.columns(4)
m1.metric("Logged", len(se))
m2.metric("Checked", len(se_chk))
m3.metric("Win rate", f"{len(se_ok)/len(se_chk):.0%}" if se_chk else "—")
m4.metric("Pending", se_pend)

# ── progress dots ──
dots = ""
for di in range(1, 6):
    day = sess["days"][str(di)]
    nt_ = len(tickers)
    nc = sum(1 for e in day["entries"].values() if e.get("was_correct") is not None)
    nl = sum(1 for e in day["entries"].values() if e.get("decision"))
    if nc == nt_ and nt_ > 0:
        dots += "🟢 "
    elif nl > 0:
        dots += "🟡 "
    else:
        dots += "⚪ "
st.markdown(f"**Progress:** {dots}")

# ── future-date hint ──
_first_log = date.fromisoformat(sess["days"]["1"]["log_date"])
if _first_log > TODAY:
    st.warning(
        f"All 5 days start **{_first_log.strftime('%a %b %d')}** — "
        f"nothing to do yet. Change the start date in a new session or wait."
    )

st.divider()


# ═══════════════════════════════════════════════════════════
# DAY CARDS — the core UI
# ═══════════════════════════════════════════════════════════
for di in range(1, 6):
    dk = str(di)
    day = sess["days"][dk]
    ld = date.fromisoformat(day["log_date"])
    cd = date.fromisoformat(day["check_date"])
    ents = day["entries"]

    n_log = sum(1 for e in ents.values() if e.get("decision"))
    n_chk = sum(1 for e in ents.values() if e.get("was_correct") is not None)
    n_t = len(tickers)
    can_log = ld <= TODAY
    can_chk = cd <= TODAY

    # day badge
    if n_chk == n_t and n_t > 0:
        n_w = sum(1 for e in ents.values() if e.get("was_correct"))
        badge = f"✅ {n_w}/{n_t} correct"
    elif n_log > 0 and can_chk:
        badge = "🔔 Ready to check T+1 price"
    elif n_log > 0:
        badge = f"📝 {n_log}/{n_t} logged · check tomorrow {cd.strftime('%a %b %d')}"
    elif can_log:
        badge = "⏳ Ready to log"
    else:
        badge = f"🔒 {ld.strftime('%a %b %d')}"

    with st.container(border=True):
        # ── header ──
        h1, h2 = st.columns([5, 3])
        today_tag = "  ← **today**" if ld == TODAY else ""
        h1.markdown(f"### Day {di} — {ld.strftime('%a %b %d')}{today_tag}")
        h1.caption(f"Log: {ld.strftime('%a %b %d')}  →  Check price: {cd.strftime('%a %b %d')}")
        h2.markdown(
            f"<div style='text-align:right;padding-top:10px'>"
            f"<b>{badge}</b></div>",
            unsafe_allow_html=True,
        )

        # ── batch buttons (compact row) ──
        unlogged = [t for t in tickers if not ents.get(t, {}).get("decision")]
        unchecked = [
            t for t in tickers
            if ents.get(t, {}).get("decision") and ents[t].get("was_correct") is None
        ]

        bb1, bb2, bb3 = st.columns([2, 2, 4])

        # Run All
        if can_log and unlogged:
            if bb1.button(
                f"🚀 Run all ({len(unlogged)})", key=f"ra_{dk}",
                use_container_width=True,
            ):
                tg = _graph()
                with st.status(f"Day {di} — running pipeline…", expanded=True) as sb:
                    for t in unlogged:
                        st.write(f"⏳ {t}…")
                        price = _price(t)
                        if price is None:
                            st.write(f"⚠️ {t}: price unavailable — skipped")
                            continue
                        try:
                            res = tg.run(t, day["log_date"], portfolio=portfolio)
                            ents[t] = _make_entry(price, res)
                            st.write(f"✅ {t}: **{ents[t]['decision']}** @ ${price:.2f}")
                        except Exception as ex:
                            st.write(f"❌ {t}: {ex}")
                    sb.update(label="Done", state="complete")
                _save(data)
                time.sleep(0.5)
                st.rerun()

        # Check All T+1
        if can_chk and unchecked:
            if bb2.button(
                f"🔍 Check T+1 ({len(unchecked)})", key=f"ca_{dk}",
                use_container_width=True,
            ):
                tg = _graph()
                with st.status(f"Day {di} — fetching T+1 prices…", expanded=True) as sb:
                    for t in unchecked:
                        e = ents[t]
                        p1 = _price(t)
                        if p1 and e.get("entry_price", 0) > 0:
                            ret = (p1 - e["entry_price"]) / e["entry_price"]
                            corr = _eval(e["decision"], ret)
                            e.update({
                                "t1_price": p1,
                                "t1_return_pct": round(ret, 6),
                                "was_correct": corr,
                                "outcome_checked_at": datetime.now().isoformat(),
                            })
                            icon = "✅" if corr else "❌"
                            st.write(f"{icon} {t}: ${p1:.2f} → {ret * 100:+.2f}%")
                            try:
                                tg.learn_from_outcome(
                                    t, day["log_date"], e["decision"],
                                    e.get("reasoning_preview", ""), ret,
                                )
                            except Exception:
                                pass
                        else:
                            st.write(f"⚠️ {t}: price unavailable")
                    sb.update(label="Done", state="complete")
                _save(data)
                time.sleep(0.5)
                st.rerun()

        # Wipe Day
        if n_log > 0:
            with bb3.popover("🗑️ Wipe day"):
                if n_chk > 0:
                    st.warning(
                        f"⚠️ {n_chk} checked outcome(s) will be removed from the log. "
                        f"Lessons already sent to KB will remain."
                    )
                st.caption(f"Delete **all** entries for Day {di}?")
                if st.button("Yes, wipe day", key=f"wd_{dk}", type="primary"):
                    for t in tickers:
                        ents[t] = {}
                    _save(data)
                    st.rerun()

        # ── per-ticker rows ──
        for t in tickers:
            e = ents.get(t, {})

            c_tick, c_info, c_act = st.columns([1, 4, 3])

            # — ticker label —
            c_tick.markdown(f"**{t}**")

            # — info column —
            with c_info:
                _ts = ""
                if e.get("logged_at"):
                    tl = _ts_label(e["logged_at"])
                    if tl:
                        _ts = f" · {tl}"

                if e.get("was_correct") is not None:
                    icon = "✅" if e["was_correct"] else "❌"
                    ret = e.get("t1_return_pct", 0)
                    st.markdown(
                        f"{icon} **{e['decision']}** {e.get('confidence', 0):.0%}  ·  "
                        f"${e.get('entry_price', 0):.2f} → ${e.get('t1_price', 0):.2f}  "
                        f"(**{ret * 100:+.2f}%**){_ts}"
                    )
                elif e.get("decision"):
                    st.markdown(
                        f"📝 **{e['decision']}** {e.get('confidence', 0):.0%}  ·  "
                        f"${e.get('entry_price', 0):.2f}{_ts}"
                        + (f"  ·  ⏳ check tomorrow {cd.strftime('%a %b %d')}" if not can_chk else "")
                    )
                else:
                    st.caption("—")

                # reasoning preview (collapsed)
                if e.get("reasoning_preview") and e.get("source", "").startswith("pipeline"):
                    with st.expander("💭 reasoning", expanded=False):
                        st.caption(e["reasoning_preview"])

            # — actions column —
            with c_act:
                a1, a2, a3, a4 = st.columns(4)

                # ── STATE: not logged ──
                if not e.get("decision"):
                    if can_log:
                        # Run single
                        if a1.button("🚀", key=f"r_{dk}_{t}", help="Run pipeline"):
                            tg = _graph()
                            with st.spinner(f"{t}…"):
                                price = _price(t)
                                if price is None:
                                    st.error(f"Could not fetch price for {t}")
                                else:
                                    try:
                                        res = tg.run(t, day["log_date"], portfolio=portfolio)
                                        ents[t] = _make_entry(price, res)
                                        _save(data)
                                        st.rerun()
                                    except Exception as ex:
                                        st.error(str(ex))

                        # Manual log
                        with a2.popover("✏️", help="Log manually"):
                            mp = st.number_input(
                                "Entry price $", min_value=0.01, step=0.01,
                                format="%.2f", key=f"lp_{dk}_{t}",
                            )
                            md = st.selectbox(
                                "Decision", ["BUY", "HOLD", "SELL"], key=f"ld_{dk}_{t}",
                            )
                            mc = st.slider(
                                "Confidence", 0.0, 1.0, 0.5, 0.05, key=f"lc_{dk}_{t}",
                            )
                            if st.button("Save", key=f"ls_{dk}_{t}", type="primary"):
                                ents[t] = {
                                    "entry_price": mp, "decision": md,
                                    "confidence": mc,
                                    "reasoning_preview": "Manual entry",
                                    "logged_at": datetime.now().isoformat(),
                                    "source": "manual",
                                    "t1_price": None, "t1_return_pct": None,
                                    "was_correct": None, "outcome_checked_at": None,
                                    "notes": "",
                                }
                                _save(data)
                                st.rerun()
                    else:
                        a1.caption("🔒")

                # ── STATE: logged, not checked ──
                elif e.get("was_correct") is None:
                    if can_chk:
                        # Auto check T+1 price
                        if a1.button("🔍", key=f"c_{dk}_{t}", help="Fetch T+1 price"):
                            tg = _graph()
                            p1 = _price(t)
                            if p1 and e.get("entry_price", 0) > 0:
                                ret = (p1 - e["entry_price"]) / e["entry_price"]
                                e.update({
                                    "t1_price": p1,
                                    "t1_return_pct": round(ret, 6),
                                    "was_correct": _eval(e["decision"], ret),
                                    "outcome_checked_at": datetime.now().isoformat(),
                                })
                                try:
                                    tg.learn_from_outcome(
                                        t, day["log_date"], e["decision"],
                                        e.get("reasoning_preview", ""), ret,
                                    )
                                except Exception:
                                    pass
                                _save(data)
                                st.rerun()
                            else:
                                st.error("Price unavailable")
                    else:
                        a1.caption("⏳")

                    # Manual T+1 price — always available
                    with a2.popover("✏️", help="Enter T+1 price manually"):
                        st.caption(f"Enter price for **{cd.isoformat()}**")
                        tp = st.number_input(
                            "T+1 price $", min_value=0.01, step=0.01,
                            format="%.2f", key=f"cp_{dk}_{t}",
                        )
                        if st.button("Save", key=f"cs_{dk}_{t}", type="primary"):
                            if e.get("entry_price", 0) > 0:
                                tg = _graph()
                                ret = (tp - e["entry_price"]) / e["entry_price"]
                                e.update({
                                    "t1_price": tp,
                                    "t1_return_pct": round(ret, 6),
                                    "was_correct": _eval(e["decision"], ret),
                                    "outcome_checked_at": datetime.now().isoformat(),
                                    "outcome_source": "manual",
                                })
                                try:
                                    tg.learn_from_outcome(
                                        t, day["log_date"], e["decision"],
                                        e.get("reasoning_preview", ""), ret,
                                    )
                                except Exception:
                                    pass
                                _save(data)
                                st.rerun()

                    # Re-run pipeline (overwrites log, clears T+1)
                    if can_log:
                        if a3.button("🔄", key=f"rr_{dk}_{t}", help="Re-run pipeline"):
                            tg = _graph()
                            with st.spinner(f"{t}…"):
                                price = _price(t)
                                if price is None:
                                    st.error("Price unavailable")
                                else:
                                    try:
                                        res = tg.run(t, day["log_date"], portfolio=portfolio)
                                        ents[t] = _make_entry(price, res, source="rerun")
                                        _save(data)
                                        st.rerun()
                                    except Exception as ex:
                                        st.error(str(ex))

                    # Delete
                    if a4.button("🗑️", key=f"d_{dk}_{t}", help="Delete entry"):
                        ents[t] = {}
                        _save(data)
                        st.rerun()

                # ── STATE: fully checked ──
                else:
                    # Re-run (via popover since KB lesson from old entry persists)
                    if can_log:
                        with a3.popover("🔄", help="Re-run pipeline"):
                            st.caption("Re-runs pipeline & replaces entry. Old KB lesson stays.")
                            if st.button("Confirm re-run", key=f"rrc_{dk}_{t}", type="primary"):
                                tg = _graph()
                                price = _price(t)
                                if price is None:
                                    st.error("Price unavailable")
                                else:
                                    try:
                                        res = tg.run(t, day["log_date"], portfolio=portfolio)
                                        ents[t] = _make_entry(price, res, source="rerun")
                                        _save(data)
                                        st.rerun()
                                    except Exception as ex:
                                        st.error(str(ex))

                    # Delete (checked — warn about KB)
                    with a4.popover("🗑️", help="Delete entry"):
                        st.caption("Entry will be removed. KB lesson from this outcome stays.")
                        if st.button("Delete", key=f"dx_{dk}_{t}", type="primary"):
                            ents[t] = {}
                            _save(data)
                            st.rerun()


# ═══════════════════════════════════════════════════════════
# KB LEARNING DELTA
# ═══════════════════════════════════════════════════════════
st.divider()

early = [e for e in se_chk if e["day"] <= 2]
late = [e for e in se_chk if e["day"] >= 4]
if early or late:
    st.markdown("#### 📈 KB learning effect")
    if sess.get("notes"):
        st.caption(f"*{sess['notes']}*")

    # grab total KB doc count for context
    _kb_n = 0
    try:
        _kb_stats = _graph().get_kb_stats()
        _kb_n = sum(_kb_stats.get(k, 0) for k in ["market_insights", "trade_history", "lessons_learned"])
    except Exception:
        pass

    lc1, lc2, lc3, lc4 = st.columns(4)
    ew = lw = None
    if early:
        ew = sum(1 for e in early if e["was_correct"]) / len(early)
        lc1.metric("Day 1–2 (KB empty)", f"{ew:.0%}")
    else:
        lc1.metric("Day 1–2", "—")
    if late:
        lw = sum(1 for e in late if e["was_correct"]) / len(late)
        lc2.metric(f"Day 4–5 (KB: {_kb_n} docs)", f"{lw:.0%}")
    else:
        lc2.metric("Day 4–5", "—")
    if ew is not None and lw is not None:
        lc3.metric("Delta Δ", f"{lw - ew:+.0%}")
    else:
        lc3.metric("Delta Δ", "—")

    if se_chk:
        ab = sum(1 for e in se_chk if e.get("t1_return_pct", 0) > 0) / len(se_chk)
        lc4.metric("Always-BUY baseline", f"{ab:.0%}")
    else:
        lc4.metric("Always-BUY baseline", "—")


# ═══════════════════════════════════════════════════════════
# ADVANCED (hidden by default)
# ═══════════════════════════════════════════════════════════
with st.expander("⚙️ Advanced — full table · downloads · KB stats"):
    if se:
        rows = []
        for e in se:
            ret = e.get("t1_return_pct")
            rows.append({
                "Day": e["day"],
                "Ticker": e["ticker"],
                "Decision": e.get("decision", ""),
                "Conf": f"{e.get('confidence', 0):.0%}",
                "Entry $": f"${e.get('entry_price', 0):.2f}",
                "Logged": _ts_label(e.get("logged_at", "")),
                "T+1 Date": e.get("check_date", ""),
                "T+1 Price $": f"${e['t1_price']:.2f}" if e.get("t1_price") else "",
                "Return": f"{ret * 100:+.2f}%" if ret is not None else "",
                "Result": (
                    "✅" if e.get("was_correct") is True
                    else ("❌" if e.get("was_correct") is False else "⏳")
                ),
            })
        st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)

    dc1, dc2 = st.columns(2)
    if os.path.exists(EXCEL_PATH):
        with open(EXCEL_PATH, "rb") as f:
            dc1.download_button(
                "📥 Excel", f.read(),
                "_eval_results.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    dc2.download_button(
        "📥 JSON",
        json.dumps(data, indent=2, default=str),
        "_eval_sessions.json",
        "application/json",
    )

    try:
        stats = _graph().get_kb_stats()
        st.caption(
            f"**KB docs:** market_insights {stats.get('market_insights', 0)}  ·  "
            f"trade_history {stats.get('trade_history', 0)}  ·  "
            f"lessons_learned {stats.get('lessons_learned', 0)}"
        )
    except Exception:
        pass
